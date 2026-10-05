"""Generate an editable HW5 Word report with screenshot placeholders."""

from __future__ import annotations

import json
import platform
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage


ROOT = Path(__file__).resolve().parent
REPORT_DIR = ROOT / "reports" / "hw05"
SCREENSHOTS = REPORT_DIR / "screenshots"
OUTPUT = REPORT_DIR / "Wang_HW5_editable.docx"


def git_value(*args: str) -> str:
    try:
        return subprocess.check_output(
            ["git", *args], cwd=ROOT, text=True, stderr=subprocess.DEVNULL
        ).strip()
    except Exception:  # noqa: BLE001
        return "unavailable"


def hardware() -> str:
    if platform.system() == "Darwin":
        try:
            text = subprocess.check_output(
                ["system_profiler", "SPHardwareDataType"], text=True, timeout=15
            )
            values = []
            for label in ("Model Name", "Chip", "Memory"):
                match = re.search(rf"^\s*{re.escape(label)}:\s*(.+)$", text, re.MULTILINE)
                if match:
                    values.append(f"{label}: {match.group(1).strip()}")
            if values:
                return ", ".join(values)
        except Exception:  # noqa: BLE001
            pass
    return f"{platform.machine()} running {platform.platform()}"


def source_lines(relative_path: str) -> list[str]:
    return (ROOT / relative_path).read_text(encoding="utf-8").splitlines()


def snippet(relative_path: str, marker: str, lines: int = 18) -> str:
    source = source_lines(relative_path)
    index = next((i for i, line in enumerate(source) if marker in line), None)
    if index is None:
        return f"# Add relevant snippet from {relative_path}: {marker}"
    start = index
    while start > 0 and source[start - 1].lstrip().startswith("@"):
        start -= 1
    return "\n".join(source[start : min(len(source), start + lines)])


def set_cell_shading(cell, fill: str) -> None:  # noqa: ANN001
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_border(cell, color: str = "D9D9D9") -> None:  # noqa: ANN001
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)
        borders.append(node)


def add_page_number(paragraph) -> None:  # noqa: ANN001
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = document.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    for style_name, size in (("Heading 1", 14), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    code = document.styles.add_style("HW5 Code", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Courier New"
    code.font.size = Pt(8)
    code.font.color.rgb = RGBColor(25, 25, 25)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.0

    footer = section.footer.paragraphs[0]
    footer.add_run("Haozhe Wang   DATA 260 HW5                                      ")
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9)


def add_hyperlink(paragraph, url: str, text: str) -> None:  # noqa: ANN001
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_code(document: Document, title: str, code: str) -> None:
    caption = document.add_paragraph()
    caption.paragraph_format.keep_with_next = True
    run = caption.add_run(title)
    run.bold = True
    run.font.size = Pt(9)
    paragraph = document.add_paragraph(code, style="HW5 Code")
    paragraph.paragraph_format.keep_with_next = True


def add_screenshot(document: Document, filename: str, caption: str) -> None:
    path = SCREENSHOTS / filename
    caption_paragraph = document.add_paragraph()
    caption_paragraph.paragraph_format.keep_with_next = True
    caption_run = caption_paragraph.add_run(caption)
    caption_run.bold = True
    caption_run.font.size = Pt(9)

    if not path.is_file():
        placeholder = document.add_paragraph()
        placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
        placeholder.paragraph_format.space_before = Pt(12)
        placeholder.paragraph_format.space_after = Pt(14)
        run = placeholder.add_run(
            f"INSERT SCREENSHOT HERE\n{filename}\nExpected source: Desktop/{filename}"
        )
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(192, 0, 0)
        return

    with PILImage.open(path) as image:
        width, height = image.size
    max_width = 7.0
    max_height = 7.0
    scale = min(max_width / width, max_height / height)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(
        str(path), width=Inches(width * scale), height=Inches(height * scale)
    )


def add_markdown_text(document: Document, path: Path) -> None:
    text = path.read_text(encoding="utf-8")
    in_code = False
    code_lines: list[str] = []
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            if in_code:
                document.add_paragraph("\n".join(code_lines), style="HW5 Code")
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("|"):
            continue
        elif line.strip():
            document.add_paragraph(line.replace("`", ""))
    if code_lines:
        document.add_paragraph("\n".join(code_lines), style="HW5 Code")


def add_config_table(document: Document) -> None:
    rows = [
        ("SID4", "8561"),
        ("PORT_BASE", "8461"),
        ("PREFIX", "s8561"),
        ("SEED", "8561"),
        ("VERIFY_SEED", "268561"),
        ("DOMAIN_ID", "1 - Clinical Trials"),
    ]
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.1)
    table.columns[1].width = Inches(4.7)
    headers = table.rows[0].cells
    headers[0].text = "Value"
    headers[1].text = "Result"
    for cell in headers:
        set_cell_shading(cell, "D9EAF7")
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].runs[0].bold = True
    for name, value in rows:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = value
        for cell in cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_api_evidence(document: Document) -> None:
    entries = [
        ("P1_Postman_02_Create_Sponsor.png", "Create Sponsor", "def create_sponsor"),
        ("P1_Postman_03_List_Sponsors.png", "List Sponsors", "def get_sponsors"),
        ("P1_Postman_04_Get_Sponsor.png", "Get Sponsor", "def get_sponsor("),
        ("P1_Postman_05_Update_Sponsor.png", "Update Sponsor", "def update_sponsor"),
        ("P1_Postman_06_Create_Trial.png", "Create Trial", "def create_trial"),
        ("P1_Postman_07_List_Trials.png", "List Trials", "def get_trials"),
        ("P1_Postman_08_Get_Trial.png", "Get Trial", "def get_trial("),
        ("P1_Postman_09_Update_Trial.png", "Update Trial", "def update_trial"),
        ("P1_Postman_10_Sponsor_Trial_Relationship.png", "Sponsor Trial Relationship", "def get_sponsor_trials"),
        ("P1_Postman_11_Delete_Sponsor_Blocked_409.png", "Delete Sponsor Blocked 409", "def delete_sponsor"),
        ("P1_Postman_12_Duplicate_NCT_409.png", "Duplicate NCT Code 409", "def create_trial"),
        ("P1_Postman_13_Invalid_NCT_422.png", "Invalid NCT Code 422", "class TrialCreate"),
        ("P1_Postman_14_Missing_Trial_404.png", "Missing Trial 404", "def find_trial_or_404"),
        ("P1_Postman_15_Delete_Trial_204.png", "Delete Trial 204", "def delete_trial"),
        ("P1_Postman_16_Delete_Sponsor_204.png", "Delete Sponsor 204", "def delete_sponsor"),
    ]
    for filename, caption, marker in entries:
        source = "code/web_application/schemas.py" if marker == "class TrialCreate" else "code/web_application/main.py"
        add_code(document, f"Relevant API code - {caption}", snippet(source, marker, 14))
        add_screenshot(document, filename, f"Postman evidence - {caption}")


def add_redux_evidence(document: Document) -> None:
    entries = [
        ("P1_Redux_17_Home_Fetch.png", "Home and Fetch", "export const fetchTrials"),
        ("P1_Redux_18_Create.png", "Create", "export const createTrial"),
        ("P1_Redux_19_Update.png", "Update", "export const updateTrial"),
        ("P1_Redux_20_Delete.png", "Delete", "export const deleteTrial"),
    ]
    for filename, caption, marker in entries:
        add_code(
            document,
            f"Redux thunk - {caption}",
            snippet("code/frontend/src/features/trials/trialsSlice.js", marker, 16),
        )
        add_screenshot(document, filename, f"Redux UI evidence - {caption}")


def add_mcp_evidence(document: Document) -> None:
    meals = [
        ("P2A_MCP_01_Search_By_Name.png", "search_meals_by_name", "async def search_meals_by_name"),
        ("P2A_MCP_02_By_Ingredient.png", "meals_by_ingredient", "async def meals_by_ingredient"),
        ("P2A_MCP_03_Random_Meal.png", "random_meal", "async def random_meal"),
        ("P2A_MCP_04_Meal_Details.png", "meal_details", "async def meal_details"),
    ]
    for filename, caption, marker in meals:
        add_code(document, f"Meals MCP tool - {caption}", snippet("code/mcp/meals_server.py", marker, 16))
        add_screenshot(document, filename, f"MCP Inspector - {caption}")

    domain = [
        ("P2B_MCP_01_Search_Success.png", "search_trials success", "def search_trials_data"),
        ("P2B_MCP_02_Search_Invalid.png", "search_trials invalid", "def search_trials_data"),
        ("P2B_MCP_03_Detail_Success.png", "get_trial_details success", "def get_trial_details_data"),
        ("P2B_MCP_04_Detail_Invalid.png", "get_trial_details invalid", "def get_trial_details_data"),
        ("P2B_MCP_05_Summary_Success.png", "summarize_sponsors success", "def summarize_sponsors_data"),
        ("P2B_MCP_06_Summary_Invalid.png", "summarize_sponsors invalid", "def summarize_sponsors_data"),
    ]
    for filename, caption, marker in domain:
        add_code(document, f"Domain MCP tool - {caption}", snippet("code/mcp/domain_server.py", marker, 17))
        add_screenshot(document, filename, f"MCP Inspector - {caption}")


def build_document() -> Document:
    document = Document()
    configure_document(document)

    document.add_paragraph("DATA 260 Homework 5", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    author = document.add_paragraph("Haozhe Wang")
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.runs[0].bold = True
    author.runs[0].font.size = Pt(14)

    remote = git_value("remote", "get-url", "origin")
    if remote.startswith("git@github.com:"):
        remote = "https://github.com/" + remote.removeprefix("git@github.com:").removesuffix(".git")
    elif remote.endswith(".git"):
        remote = remote[:-4]
    repo_paragraph = document.add_paragraph()
    repo_paragraph.add_run("GitHub repository: ").bold = True
    add_hyperlink(repo_paragraph, remote, remote)
    document.add_paragraph(f"Commit used for report: {git_value('rev-parse', 'HEAD')}")
    document.add_paragraph("Submission tag: hw5")
    document.add_paragraph(f"Hardware: {hardware()}")
    document.add_paragraph("Local model: qwen3:8b through Ollama")

    document.add_heading("Section 0 Personal Configuration", level=1)
    add_config_table(document)
    document.add_page_break()

    document.add_heading("Part 1 Database and API", level=1)
    document.add_paragraph(
        "The FastAPI and MySQL service adds Sponsor and Trial CRUD, a Sponsor to Trial "
        "relationship, unique NCT code validation, pagination, timestamps, and explicit "
        "404, 409, and 422 responses."
    )
    add_code(document, "SQLAlchemy relationship models", snippet("code/web_application/models.py", "class Sponsor", 48))
    add_screenshot(document, "P1_DB_Relationship.png", "MySQL relationship evidence")
    add_api_evidence(document)

    document.add_page_break()
    document.add_heading("Part 1 Redux Toolkit Client", level=1)
    document.add_paragraph(
        "Redux Toolkit owns the Trial state. Axios backed async thunks fetch, create, "
        "update, and delete records, and fulfilled actions update the Home screen."
    )
    add_redux_evidence(document)

    document.add_page_break()
    document.add_heading("Part 2 MCP Tool Servers", level=1)
    document.add_paragraph(
        "The Meals STDIO server exposes four TheMealDB tools. The domain STDIO server "
        "exposes exactly three clinical trial tools with the shared ok data error envelope."
    )
    add_mcp_evidence(document)

    document.add_page_break()
    document.add_heading("Part 3 Tool Contracts Under Stress", level=1)
    add_markdown_text(document, REPORT_DIR / "PART3_TOOL_CONTRACTS.md")
    add_code(document, "Bounded timeout and exponential backoff", snippet("code/mcp/retry_policy.py", "def run_with_retry", 34))
    add_screenshot(document, "P3_Retry_01_Demo.png", "Retry policy demonstration")
    add_markdown_text(document, REPORT_DIR / "PART3_RETRY_ANALYSIS.md")
    add_screenshot(document, "P3_Retry_02_Experiment.png", "Seeded 150 call experiment")

    document.add_page_break()
    document.add_heading("Part 4 Safe Tool Entry Point and Offline Tests", level=1)
    add_code(document, "Single JSON string tool boundary", snippet("code/mcp/tool_executor.py", "def execute_tool", 42))
    add_markdown_text(document, REPORT_DIR / "PART4_OFFLINE_TESTS.md")
    add_screenshot(document, "P4_Execute_Tool_Tests.png", "Offline test runner")

    document.add_page_break()
    document.add_heading("Part 5 Agent Loop Safety Metrics and Reflection", level=1)
    add_code(document, "Domain safety rule", snippet("code/mcp/tool_executor.py", "if clean_name == \"summarize_sponsors\"", 24))
    add_screenshot(document, "P5_Safety_Demo.png", "Allowed and blocked safety calls")
    add_code(document, "Bounded Ollama agent loop", snippet("code/mcp/agent_loop.py", "def run_agent", 56))
    add_markdown_text(document, REPORT_DIR / "METRICS.md")
    add_screenshot(document, "P5_Agent_Scenarios.png", "Four local Ollama scenarios")
    add_markdown_text(document, REPORT_DIR / "REFLECTION.md")

    document.add_page_break()
    document.add_heading("AI Use", level=1)
    add_markdown_text(document, REPORT_DIR / "AI_USE.md")
    document.add_heading("Reproducibility and Raw Evidence", level=1)
    document.add_paragraph(
        "The repository includes RUN_LOG.txt, verification.json, all 150 fault injection "
        "records, machine readable MCP calls, the Part 5 JSONL log, metrics, and the "
        "labeled screenshots used in this report."
    )
    return document


def main() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.core_properties.title = "DATA 260 Homework 5"
    document.core_properties.author = "Haozhe Wang"
    document.save(OUTPUT)

    missing = []
    for paragraph in document.paragraphs:
        if "INSERT SCREENSHOT HERE" in paragraph.text:
            missing.append(paragraph.text.splitlines()[1])
    print(f"Generated editable Word report: {OUTPUT}")
    print(f"Missing screenshot placeholders: {len(missing)}")
    for filename in missing:
        print(f"- {filename}")


if __name__ == "__main__":
    main()
